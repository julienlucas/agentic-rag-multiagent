"""
Formats acceptés et leur conversion en « pages » de markdown.

Chaque format est ramené à une liste de pages : le PageStore et les outils de l'agent
(read_page, grep, navigate) ne voient jamais la différence entre un PDF et un tableur.

| Famille  | Extensions                         | Traitement                                               |
|----------|------------------------------------|----------------------------------------------------------|
| pdf      | .pdf                               | Mistral OCR                                              |
| office   | .docx .pptx .odt .odp .rtf .doc .ppt | LibreOffice → PDF (visualiseur à la page) puis OCR ;     |
|          |                                    | sans LibreOffice, DOCX / PPTX partent tels quels à l'OCR |
| image    | .png .jpg .jpeg .webp              | Mistral OCR (image_url) : une image = une page           |
| text     | .txt .md .markdown                 | Lu directement, découpé en pages de ~3 000 caractères    |
| html     | .html .htm                         | HTML → markdown, puis comme du texte                     |
| sheet    | .xlsx .csv                         | Un tableau markdown par tranche de 50 lignes et par feuille |

Le type se vérifie aussi par le contenu (signature) : une extension trompeuse est refusée.
"""
import csv
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from typing import Callable, List, Optional

KINDS = {
    ".pdf": "pdf",
    ".docx": "office", ".pptx": "office", ".odt": "office", ".odp": "office", ".rtf": "office",
    ".doc": "office", ".ppt": "office",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".webp": "image",
    ".txt": "text", ".md": "text", ".markdown": "text",
    ".html": "html", ".htm": "html",
    ".xlsx": "sheet", ".csv": "sheet",
}
EXTENSIONS = sorted(KINDS)

MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
}

TEXT_PAGE_CHARS = 3000
SHEET_PAGE_ROWS = 50
MAX_PAGES = 500


class UnsupportedFormat(ValueError):
    pass


def suffix(name: str) -> str:
    return os.path.splitext(name or "")[1].lower()


def kind_of(name: str) -> Optional[str]:
    return KINDS.get(suffix(name))


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def check(name: str, content: bytes) -> str:
    """La famille du fichier, si son extension est acceptée ET que son contenu la confirme."""
    ext = suffix(name)
    kind = KINDS.get(ext)
    if kind is None:
        raise UnsupportedFormat(f"Format non pris en charge : {ext or 'sans extension'}. "
                                f"Formats acceptés : {', '.join(EXTENSIONS)}.")
    head = content[:8]
    ok = {
        ".pdf": head.startswith(b"%PDF"),
        ".png": head.startswith(b"\x89PNG"),
        ".jpg": head.startswith(b"\xff\xd8"),
        ".jpeg": head.startswith(b"\xff\xd8"),
        ".webp": content[:4] == b"RIFF" and content[8:12] == b"WEBP",
        ".docx": head.startswith(b"PK"), ".pptx": head.startswith(b"PK"), ".xlsx": head.startswith(b"PK"),
        ".odt": head.startswith(b"PK"), ".odp": head.startswith(b"PK"),
        ".doc": head.startswith(b"\xd0\xcf\x11\xe0"), ".ppt": head.startswith(b"\xd0\xcf\x11\xe0"),
        ".rtf": head.startswith(b"{\\rtf"),
    }.get(ext, b"\x00" not in content[:4096])  # texte, HTML, CSV : pas d'octet nul
    if not ok:
        raise UnsupportedFormat(f"Le contenu de {os.path.basename(name)} ne correspond pas à son extension.")
    return kind


# ── Conversion en PDF (visualiseur) ─────────────────────────────────────────────

def soffice_available(binary: str) -> bool:
    return shutil.which(binary) is not None


def convert_to_pdf(name: str, content: bytes, binary: str, timeout: int = 180) -> Optional[bytes]:
    """DOCX, PPTX… → PDF par LibreOffice sans interface. None si LibreOffice manque ou échoue."""
    if not soffice_available(binary):
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, f"source{suffix(name)}")
        with open(src, "wb") as f:
            f.write(content)
        try:
            subprocess.run(
                [binary, "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                check=True, timeout=timeout, capture_output=True,
                env={**os.environ, "HOME": tmp},  # profil LibreOffice jetable, pas de verrou partagé
            )
        except (subprocess.SubprocessError, OSError):
            return None
        out = os.path.join(tmp, "source.pdf")
        return open(out, "rb").read() if os.path.exists(out) else None


# ── Nombre de pages (facturation à la réception) ────────────────────────────────

def pdf_page_count(content: bytes) -> int:
    from pypdf import PdfReader
    return len(PdfReader(io.BytesIO(content)).pages)


def estimate_pages(name: str, content: bytes) -> int:
    """
    Pages facturées à la réception. Exact pour PDF, image, texte, tableur et PPTX ; estimé pour
    les autres formats bureautiques (~450 mots par page), réajusté après l'indexation.
    """
    kind, ext = kind_of(name), suffix(name)
    if kind == "pdf":
        return max(1, pdf_page_count(content))
    if kind == "image":
        return 1
    if kind in ("text", "html"):
        return len(text_pages(_text_of(ext, content)))
    if kind == "sheet":
        return len(sheet_pages(name, content))
    if ext == ".pptx":
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            return max(1, sum(1 for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)))
    if ext in (".docx", ".odt"):
        inner = "word/document.xml" if ext == ".docx" else "content.xml"
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            xml = z.read(inner).decode("utf-8", errors="ignore")
        words = len(re.sub(r"<[^>]+>", " ", xml).split())
        return max(1, -(-words // 450))
    return max(1, -(-len(content) // 30_000))  # .doc, .ppt, .rtf, .odp : ordre de grandeur


# ── Extraction des pages ────────────────────────────────────────────────────────

def _text_of(ext: str, content: bytes) -> str:
    text = _decode(content)
    if ext in (".html", ".htm"):
        from bs4 import BeautifulSoup
        from markdownify import MarkdownConverter
        soup = BeautifulSoup(text, "html.parser")
        for tag in soup(["script", "style", "noscript", "template"]):
            tag.decompose()  # leur contenu n'est pas du texte lisible
        text = MarkdownConverter(heading_style="ATX").convert_soup(soup)
    return text


def text_pages(text: str, size: int = TEXT_PAGE_CHARS) -> List[str]:
    """Pages de ~size caractères, coupées entre deux paragraphes quand c'est possible."""
    text = text.replace("\r\n", "\n").strip()
    if not text:
        return []
    pages, current = [], ""
    for para in re.split(r"\n\s*\n", text):
        while len(para) > size:  # paragraphe géant : coupé à la ligne ou au mot
            cut = max(para.rfind("\n", 0, size), para.rfind(" ", 0, size))
            cut = cut if cut > size // 2 else size
            if current:
                pages.append(current)
                current = ""
            pages.append(para[:cut].strip())
            para = para[cut:].strip()
        if current and len(current) + len(para) + 2 > size:
            pages.append(current)
            current = para
        else:
            current = f"{current}\n\n{para}" if current else para
    if current:
        pages.append(current)
    return pages


def _md_cell(value) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def _table_pages(title: str, rows: List[List]) -> List[str]:
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    header, body = rows[0], rows[1:] or [[]]
    head = "| " + " | ".join(_md_cell(c) for c in list(header) + [""] * (width - len(header))) + " |"
    sep = "|" + "---|" * width
    pages = []
    for start in range(0, len(body), SHEET_PAGE_ROWS):
        part = body[start:start + SHEET_PAGE_ROWS]
        lines = ["| " + " | ".join(_md_cell(c) for c in list(r) + [""] * (width - len(r))) + " |" for r in part]
        span = f"lignes {start + 2} à {start + 1 + len(part)}"
        pages.append(f"## {title} ({span})\n\n" + "\n".join([head, sep, *lines]))
    return pages


def sheet_pages(name: str, content: bytes) -> List[str]:
    if suffix(name) == ".csv":
        text = _decode(content)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return _table_pages(os.path.basename(name), list(csv.reader(io.StringIO(text), dialect)))
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    pages = []
    for ws in wb.worksheets:
        pages += _table_pages(f"Feuille {ws.title}", [list(r) for r in ws.iter_rows(values_only=True)])
    return pages


def extract_pages(name: str, content: bytes, ocr: Callable[[bytes, str], List[str]],
                  pdf: Optional[bytes] = None) -> List[str]:
    """
    Le markdown de chaque page. `ocr(content, mime)` appelle Mistral OCR ; `pdf` est la conversion
    PDF d'un fichier bureautique, quand LibreOffice l'a produite (elle est alors OCRisée à la place).
    """
    kind, ext = kind_of(name), suffix(name)
    if kind == "pdf":
        pages = ocr(content, MIME[".pdf"])
    elif kind == "office":
        if pdf is not None:
            pages = ocr(pdf, MIME[".pdf"])
        elif ext in (".docx", ".pptx"):
            pages = ocr(content, MIME[ext])
        else:
            raise UnsupportedFormat(f"{ext} demande LibreOffice, absent de ce serveur.")
    elif kind == "image":
        pages = ocr(content, MIME[ext])
    elif kind in ("text", "html"):
        pages = text_pages(_text_of(ext, content))
    elif kind == "sheet":
        pages = sheet_pages(name, content)
    else:
        raise UnsupportedFormat(f"Format non pris en charge : {ext}")
    return pages[:MAX_PAGES]
